"""
Office document parser for VibeMask.
Supports: .docx, .xlsx, .pptx
"""

from pathlib import Path
from typing import List, Tuple, Optional
from dataclasses import dataclass


@dataclass
class TextBlock:
    """A block of text extracted from a document."""
    text: str
    location: str  # e.g., "paragraph:3", "cell:A1", "slide:2"
    

def read_docx(file_path: Path) -> Tuple[List[TextBlock], any]:
    """
    Read a Word document and extract all text blocks.
    
    Args:
        file_path: Path to .docx file
        
    Returns:
        Tuple of (text_blocks, document_object for later saving)
    """
    from docx import Document
    
    doc = Document(file_path)
    blocks = []
    
    # Extract paragraphs
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            blocks.append(TextBlock(
                text=para.text,
                location=f"paragraph:{i}"
            ))
    
    # Extract tables
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if cell.text.strip():
                    blocks.append(TextBlock(
                        text=cell.text,
                        location=f"table:{t_idx}:row:{r_idx}:col:{c_idx}"
                    ))
    
    # Extract headers and footers
    for section in doc.sections:
        if section.header:
            for para in section.header.paragraphs:
                if para.text.strip():
                    blocks.append(TextBlock(
                        text=para.text,
                        location="header"
                    ))
        if section.footer:
            for para in section.footer.paragraphs:
                if para.text.strip():
                    blocks.append(TextBlock(
                        text=para.text,
                        location="footer"
                    ))
    
    return blocks, doc


def write_docx(doc: any, file_path: Path, replacements: dict):
    """
    Write back a Word document with replacements applied.
    
    Args:
        doc: Document object from read_docx
        file_path: Output path
        replacements: Dict of {original -> masked}
    """
    # Replace in paragraphs
    for para in doc.paragraphs:
        for original, masked in replacements.items():
            if original in para.text:
                # Need to handle runs properly
                for run in para.runs:
                    if original in run.text:
                        run.text = run.text.replace(original, masked)
    
    # Replace in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for original, masked in replacements.items():
                    if original in cell.text:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                if original in run.text:
                                    run.text = run.text.replace(original, masked)
    
    # Replace in headers/footers
    for section in doc.sections:
        if section.header:
            for para in section.header.paragraphs:
                for run in para.runs:
                    for original, masked in replacements.items():
                        if original in run.text:
                            run.text = run.text.replace(original, masked)
        if section.footer:
            for para in section.footer.paragraphs:
                for run in para.runs:
                    for original, masked in replacements.items():
                        if original in run.text:
                            run.text = run.text.replace(original, masked)
    
    doc.save(file_path)


def read_xlsx(file_path: Path) -> Tuple[List[TextBlock], any]:
    """
    Read an Excel workbook and extract all cell values.
    """
    from openpyxl import load_workbook
    
    wb = load_workbook(file_path)
    blocks = []
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.strip():
                    blocks.append(TextBlock(
                        text=str(cell.value),
                        location=f"sheet:{sheet_name}:cell:{cell.coordinate}"
                    ))
    
    return blocks, wb


def write_xlsx(wb: any, file_path: Path, replacements: dict):
    """
    Write back an Excel workbook with replacements applied.
    
    IMPORTANT: Sort by key length descending to prevent substring conflicts.
    E.g., "赵乙丑"->"崔兆鹏" must be replaced before "赵乙"->"卜凯"
    """
    # Sort replacements by key length (longest first)
    sorted_replacements = sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True)
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    new_value = cell.value
                    for original, masked in sorted_replacements:
                        if original in new_value:
                            new_value = new_value.replace(original, masked)
                    cell.value = new_value
    
    wb.save(file_path)


def read_pptx(file_path: Path) -> Tuple[List[TextBlock], any]:
    """
    Read a PowerPoint presentation and extract all text.
    """
    from pptx import Presentation
    
    prs = Presentation(file_path)
    blocks = []
    
    for slide_idx, slide in enumerate(prs.slides):
        for shape_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                for para_idx, para in enumerate(shape.text_frame.paragraphs):
                    text = para.text.strip()
                    if text:
                        blocks.append(TextBlock(
                            text=text,
                            location=f"slide:{slide_idx}:shape:{shape_idx}:para:{para_idx}"
                        ))
            
            # Handle tables in slides
            if shape.has_table:
                table = shape.table
                for row_idx, row in enumerate(table.rows):
                    for col_idx, cell in enumerate(row.cells):
                        if cell.text.strip():
                            blocks.append(TextBlock(
                                text=cell.text,
                                location=f"slide:{slide_idx}:table:row:{row_idx}:col:{col_idx}"
                            ))
    
    return blocks, prs


def write_pptx(prs: any, file_path: Path, replacements: dict):
    """
    Write back a PowerPoint presentation with replacements applied.
    """
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        for original, masked in replacements.items():
                            if original in run.text:
                                run.text = run.text.replace(original, masked)
            
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                for original, masked in replacements.items():
                                    if original in run.text:
                                        run.text = run.text.replace(original, masked)
    
    prs.save(file_path)


def get_combined_text(blocks: List[TextBlock]) -> str:
    """Combine all text blocks into a single string for detection."""
    return "\n".join(block.text for block in blocks)


def is_office_file(file_path: Path) -> bool:
    """Check if file is an Office document."""
    return file_path.suffix.lower() in {'.docx', '.xlsx', '.pptx'}


def read_office_file(file_path: Path) -> Tuple[List[TextBlock], any]:
    """
    Read any supported Office file.
    
    Returns:
        Tuple of (text_blocks, document_object)
    """
    suffix = file_path.suffix.lower()
    
    if suffix == '.docx':
        return read_docx(file_path)
    elif suffix == '.xlsx':
        return read_xlsx(file_path)
    elif suffix == '.pptx':
        return read_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")


def write_office_file(doc: any, file_path: Path, replacements: dict, file_type: str):
    """
    Write back any supported Office file.
    """
    if file_type == '.docx':
        write_docx(doc, file_path, replacements)
    elif file_type == '.xlsx':
        write_xlsx(doc, file_path, replacements)
    elif file_type == '.pptx':
        write_pptx(doc, file_path, replacements)
    else:
        raise ValueError(f"Unsupported file format: {file_type}")


def mask_office_file(
    input_path: Path,
    output_path: Optional[Path] = None,
    dry_run: bool = False
) -> dict:
    """
    Mask sensitive information in an Office file.
    
    Args:
        input_path: Path to input file
        output_path: Path for output (default: overwrites input)
        dry_run: If True, only return detection results without writing
        
    Returns:
        Dict with masking results
    """
    from ..core.merger import merge_spans
    from ..detector.regex import detect_by_regex
    from ..detector.heuristic import detect_chinese_names
    from ..masker.placeholder import PlaceholderGenerator
    from ..vault.storage import VaultStorage
    
    # Read document
    blocks, doc = read_office_file(input_path)
    combined_text = get_combined_text(blocks)
    
    # Detect entities
    spans = []
    spans.extend(detect_by_regex(combined_text))
    spans.extend(detect_chinese_names(combined_text))
    
    # Merge
    final_spans = merge_spans(spans, combined_text)
    
    if not final_spans:
        return {
            "success": True,
            "file": str(input_path),
            "entities_found": 0,
            "stats": {},
        }
    
    # Generate replacements
    generator = PlaceholderGenerator()
    replacements = {}
    stats = {}
    
    for span in final_spans:
        masked = generator.generate(span.text, span.type)
        replacements[span.text] = masked
        stats[span.type.value] = stats.get(span.type.value, 0) + 1
    
    result = {
        "success": True,
        "file": str(input_path),
        "entities_found": len(final_spans),
        "stats": stats,
        "replacements": replacements if dry_run else None,
    }
    
    if dry_run:
        return result
    
    # Apply replacements and save
    if output_path is None:
        output_path = input_path
    
    vault = VaultStorage(str(input_path.parent))
    with vault.batch("mask-session"):
        for span in final_spans:
            proposed = replacements[span.text]
            replacements[span.text] = vault.get_or_create_mapping(
                original=span.text,
                entity_type=span.type.value,
                masked=proposed,
                source=span.source.value,
                confidence=span.confidence,
            )

        write_office_file(doc, output_path, replacements, input_path.suffix.lower())

        mappings = {v: k for k, v in replacements.items()}
        session_id = vault.create_session(
            [str(input_path)], mappings, stats, output_files=[str(output_path)]
        )
    
    result["session_id"] = session_id
    result["output"] = str(output_path)
    
    return result


def restore_office_file(
    input_path: Path,
    session_id: Optional[str] = None,
    output_path: Optional[Path] = None
) -> dict:
    """
    Restore masked Office file to original.
    """
    from ..vault.storage import VaultStorage
    
    vault = VaultStorage(str(input_path.parent))
    
    if session_id:
        session = vault.get_session(session_id)
        if not session:
            return {"success": False, "error": "Session not found"}
        mappings = session.mappings
    else:
        mappings = vault.get_all_mappings()
    
    # Read document
    blocks, doc = read_office_file(input_path)
    
    # Reverse mappings: {masked -> original}
    # mappings is already in this format from session
    
    if output_path is None:
        output_path = input_path
    
    # Apply restoration
    write_office_file(doc, output_path, mappings, input_path.suffix.lower())
    
    return {
        "success": True,
        "file": str(input_path),
        "output": str(output_path),
    }
