import os
import sys
import zipfile
from pathlib import Path

import pytest

# Ensure the repository root is on sys.path when pytest uses importlib mode.
REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


@pytest.mark.skipif(not _docx_available(), reason="python-docx not installed")
def test_docx_cross_run_mask_restore_roundtrip(tmp_path: Path):
    # Create a DOCX where a name is split across runs (<w:t> nodes).
    from docx import Document

    input_path = tmp_path / "input.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("张")
    p.add_run("三")
    doc.save(input_path)

    from vibemask.core import factory

    # Mask with a replacement that spans runs.
    replacements = {"张三": "{{PERSON_000001}}"}
    masked_path = tmp_path / "input_masked.docx"
    proc = factory.get_processor(input_path)
    proc.load()
    proc.extract_segments()
    assert proc.replace_text(replacements) == 1
    proc.save(masked_path)

    # Ensure masked file is a valid docx zip.
    with zipfile.ZipFile(masked_path) as zf:
        assert "word/document.xml" in zf.namelist()

    # Restore and validate text roundtrip.
    restored_path = tmp_path / "input_restored.docx"
    proc2 = factory.get_processor(masked_path)
    proc2.load()
    proc2.extract_segments()
    assert proc2.replace_text({"{{PERSON_000001}}": "张三"}) == 1
    proc2.save(restored_path)

    restored_text = factory.extract_text(restored_path)
    assert "张三" in restored_text
    assert "{{PERSON_000001}}" not in restored_text

    # Open with python-docx to ensure Office format isn't corrupted.
    Document(restored_path)


def test_placeholder_phone_preserves_separators():
    from vibemask.masker.placeholder import PlaceholderGenerator
    from vibemask.core.span import EntityType

    g = PlaceholderGenerator()
    original = "138-1234-5678"
    masked = g.generate(original, EntityType.PHONE)
    assert len(masked) == len(original)
    assert masked[3] == "-"
    assert masked[8] == "-"
    assert masked.replace("-", "").isdigit()


def test_vault_enforces_unique_masked_values(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    # Isolate vault into the test temp HOME so we don't touch real user data.
    monkeypatch.setenv("HOME", str(tmp_path))

    from vibemask.vault.storage import VaultStorage

    vault = VaultStorage(str(tmp_path / "project"))

    m1 = vault.get_or_create_mapping(original="张三", entity_type="PERSON", masked="{{PERSON_000001}}")
    m2 = vault.get_or_create_mapping(original="李四", entity_type="PERSON", masked="{{PERSON_000001}}")
    assert m1 != m2

    # Ensure reverse lookup stays unambiguous.
    assert vault.get_mapping_by_masked(m1) == "张三"
    assert vault.get_mapping_by_masked(m2) == "李四"


def test_vault_handles_many_short_cjk_mask_collisions(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("HOME", str(tmp_path))

    from vibemask.vault.storage import VaultStorage

    vault = VaultStorage(str(tmp_path / "project"))

    masked_values = {
        vault.get_or_create_mapping(
            original=f"姓名{i:03d}",
            entity_type="PERSON",
            masked="{{PERSON_000001}}",
        )
        for i in range(120)
    }

    assert len(masked_values) == 120


@pytest.mark.skipif(os.name == "nt", reason="POSIX permission bits only")
def test_vault_uses_private_filesystem_permissions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setenv("HOME", str(tmp_path))

    from vibemask.vault.storage import VaultStorage

    vault = VaultStorage(str(tmp_path / "project"))

    assert (vault.vault_path.parent.stat().st_mode & 0o777) == 0o700
    assert (vault.vault_path.stat().st_mode & 0o777) == 0o600
