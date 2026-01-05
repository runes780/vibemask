import sys
import zipfile
from pathlib import Path

import pytest

# Ensure the repository root is on sys.path when pytest uses importlib mode.
REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


def _openpyxl_available() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except Exception:
        return False


@pytest.mark.skipif(not _docx_available(), reason="python-docx not installed")
def test_docx_sample_roundtrip_preserves_text(tmp_path: Path):
    from vibemask.core import factory
    from docx import Document

    sample = Path(__file__).parent / "7931b76a183e4a54976b10cbf2d2fc1c.docx"
    assert sample.exists()

    original_copy = tmp_path / sample.name
    original_copy.write_bytes(sample.read_bytes())

    original_text = factory.extract_text(original_copy)
    needle = "发改办高技〔2016〕937号"
    assert needle in original_text

    masked_path = tmp_path / "sample_masked.docx"
    restored_path = tmp_path / "sample_restored.docx"

    proc = factory.get_processor(original_copy)
    proc.load()
    proc.extract_segments()
    assert proc.replace_text({needle: "X" * len(needle)}) == 1
    proc.save(masked_path)

    with zipfile.ZipFile(masked_path) as zf:
        assert "word/document.xml" in zf.namelist()

    proc2 = factory.get_processor(masked_path)
    proc2.load()
    proc2.extract_segments()
    assert proc2.replace_text({"X" * len(needle): needle}) == 1
    proc2.save(restored_path)

    restored_text = factory.extract_text(restored_path)
    assert restored_text == original_text

    # Ensure Office can parse it (python-docx open).
    Document(restored_path)


@pytest.mark.skipif(not _openpyxl_available(), reason="openpyxl not installed")
def test_xlsx_sample_roundtrip_preserves_text(tmp_path: Path):
    from vibemask.core import factory
    from openpyxl import load_workbook

    sample = Path(__file__).parent / "b1868fc7-e4dc-4415-a85d-37a9d714dec7.xlsx"
    assert sample.exists()

    original_copy = tmp_path / sample.name
    original_copy.write_bytes(sample.read_bytes())

    original_text = factory.extract_text(original_copy)
    needle = "蔡登"
    assert needle in original_text

    masked_path = tmp_path / "sample_masked.xlsx"
    restored_path = tmp_path / "sample_restored.xlsx"

    proc = factory.get_processor(original_copy)
    proc.load()
    proc.extract_segments()
    assert proc.replace_text({needle: "赵甲"}) == 1
    proc.save(masked_path)

    proc2 = factory.get_processor(masked_path)
    proc2.load()
    proc2.extract_segments()
    assert proc2.replace_text({"赵甲": needle}) == 1
    proc2.save(restored_path)

    restored_text = factory.extract_text(restored_path)
    assert restored_text == original_text

    # Ensure Excel can parse it (openpyxl load).
    load_workbook(restored_path)

